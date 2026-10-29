from __future__ import annotations

import re
from datetime import datetime
from datetime import timezone
from pathlib import Path
from typing import Any

from data.sec_normalizer import build_financial_snapshot
from data.sec_tools import SecClientError
from data.sec_tools import fetch_company_facts
from data.sec_tools import fetch_company_submissions
from data.sec_tools import lookup_company
from valuation.schemas import ResolvedCompany

from .schemas import CompsCompanyRecord
from .schemas import CompsExcelArtifact
from .schemas import PeerCandidate
from .schemas import PeerProposal
from .schemas import PeerReviewDecision


MIME_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def default_export_dir() -> Path:
    return Path(__file__).resolve().parents[2] / "api" / ".adk" / "exports"


def _safe_name(value: str) -> str:
    cleaned = re.sub(r"[^A-Za-z0-9_.-]+", "_", value.strip())
    return cleaned.strip("._") or "comps"


def _millions(value: float | None) -> float | None:
    return None if value is None else value * 1000.0


def _lookup_query(peer: PeerCandidate) -> str:
    return peer.ticker or peer.company_name


def _resolved_company(query: str) -> ResolvedCompany:
    result = lookup_company(query)
    if result.get("status") != "success":
        raise ValueError(result.get("message") or f"No SEC match found for {query}.")
    best = result["best_match"]
    return ResolvedCompany(
        query=query,
        title=best["title"],
        ticker=best["ticker"],
        cik=best["cik"],
        cik_padded=best["cik_padded"],
        data_source="sec",
        matches=result.get("matches", []),
    )


def build_sec_company_record(
    *,
    peer: PeerCandidate,
    peer_category: str,
) -> CompsCompanyRecord:
    query = _lookup_query(peer)
    try:
        company = _resolved_company(query)
        if company.cik is None:
            raise ValueError(f"SEC CIK is missing for {query}.")
        facts = fetch_company_facts(company.cik)
        submissions = fetch_company_submissions(company.cik)
        snapshot = build_financial_snapshot(company, facts, submissions)
    except (SecClientError, ValueError) as exc:
        return CompsCompanyRecord(
            query=query,
            company_name=peer.company_name,
            ticker=peer.ticker,
            peer_category=peer_category,  # type: ignore[arg-type]
            status="not_found" if isinstance(exc, ValueError) else "failed",
            limitations=[str(exc)],
        )

    return CompsCompanyRecord(
        query=query,
        company_name=snapshot.company_name,
        ticker=snapshot.ticker,
        cik=snapshot.cik,
        peer_category=peer_category,  # type: ignore[arg-type]
        status=snapshot.status,
        fiscal_year=snapshot.fiscal_year,
        filing_date=snapshot.filing_date,
        revenue_mm=_millions(snapshot.revenue_base),
        gross_profit_mm=_millions(snapshot.gross_profit),
        operating_income_mm=_millions(snapshot.operating_income),
        depreciation_amortization_mm=_millions(snapshot.depreciation_amortization),
        cash_mm=_millions(snapshot.cash_and_non_operating_asset),
        debt_mm=_millions(snapshot.debt_value_book),
        shares_outstanding_mm=_millions(snapshot.shares_outstanding),
        sources=[source.model_dump(mode="json") for source in snapshot.sources],
        limitations=snapshot.notes,
    )


def build_sec_comps_records(
    *,
    target_company: str,
    target_ticker: str | None,
    selected_peers: list[PeerCandidate],
) -> list[CompsCompanyRecord]:
    target = PeerCandidate(
        company_name=target_company,
        ticker=target_ticker,
        peer_category="core",
        rationale="Target company.",
    )
    records = [
        build_sec_company_record(peer=target, peer_category="target"),
    ]
    records.extend(
        build_sec_company_record(peer=peer, peer_category=peer.peer_category)
        for peer in selected_peers
    )
    return records


def _source_for_metric(record: CompsCompanyRecord, metric: str) -> dict[str, Any] | None:
    for source in record.sources:
        if source.get("metric") == metric:
            return source
    return None


def _comment_text(source: dict[str, Any] | None, fallback: str) -> str:
    if source is None:
        return fallback
    parts = [
        source.get("source"),
        f"metric={source.get('metric')}",
        f"form={source.get('form')}",
        f"fy={source.get('fiscal_year')}",
        f"filed={source.get('filing_date')}",
        f"accn={source.get('accession_number')}",
        f"xbrl={source.get('xbrl_tag')}",
    ]
    if source.get("notes"):
        parts.append(str(source["notes"]))
    return "; ".join(str(part) for part in parts if part)


def create_comps_workbook(
    *,
    workflow_request: dict[str, Any],
    proposal: PeerProposal,
    review: PeerReviewDecision,
    records: list[CompsCompanyRecord],
    export_dir: Path | None = None,
) -> CompsExcelArtifact:
    try:
        from openpyxl import Workbook
        from openpyxl.comments import Comment
        from openpyxl.styles import Alignment
        from openpyxl.styles import Font
        from openpyxl.styles import PatternFill
    except ModuleNotFoundError as exc:
        raise RuntimeError(
            "openpyxl is required to generate comparables Excel workbooks. "
            "Install project dependencies before resuming peer approval."
        ) from exc

    export_root = export_dir or default_export_dir()
    export_root.mkdir(parents=True, exist_ok=True)
    job_id = str(workflow_request.get("job_id") or workflow_request.get("request_id") or "comps")
    filename = f"{_safe_name(job_id)}_comps.xlsx"
    path = export_root / filename

    wb = Workbook()
    ws = wb.active
    ws.title = "Comps Analysis"
    sources_ws = wb.create_sheet("Sources")
    review_ws = wb.create_sheet("Peer Review")

    dark_fill = PatternFill("solid", fgColor="17365D")
    header_fill = PatternFill("solid", fgColor="D9E2F3")
    stat_fill = PatternFill("solid", fgColor="F2F2F2")
    input_font = Font(name="Times New Roman", size=11, color="0000FF")
    formula_font = Font(name="Times New Roman", size=11, color="000000")
    header_font = Font(name="Times New Roman", size=11, bold=True)
    section_font = Font(name="Times New Roman", size=11, bold=True, color="FFFFFF")

    title = f"{proposal.target_company} Comparable Company Analysis"
    ws["A1"] = title
    ws["A2"] = "Peer-reviewed SEC/XBRL historical trading comps input file"
    ws["A3"] = (
        "All financial statement figures in USD millions. Market data cells are blank "
        "unless supplied by an approved market data source."
    )
    for cell in ("A1", "A2", "A3"):
        ws[cell].font = Font(name="Times New Roman", size=12 if cell != "A3" else 11, bold=cell == "A1")
    ws.merge_cells("A1:W1")
    ws.merge_cells("A2:W2")
    ws.merge_cells("A3:W3")

    headers = [
        "Company",
        "Ticker",
        "Peer Category",
        "CIK",
        "Fiscal Year",
        "Filing Date",
        "Status",
        "Revenue",
        "Gross Profit",
        "Gross Margin",
        "Operating Income",
        "D&A",
        "EBITDA",
        "EBITDA Margin",
        "Cash",
        "Debt",
        "Net Debt",
        "Shares Out.",
        "Share Price",
        "Market Cap",
        "Enterprise Value",
        "EV/Revenue",
        "EV/EBITDA",
    ]
    header_row = 5
    for col_index, header in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=col_index, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    metric_sources = {
        8: "revenue_base",
        9: "gross_profit",
        11: "operating_income",
        12: "depreciation_amortization",
        15: "cash_and_non_operating_asset",
        16: "debt_value_book",
        18: "shares_outstanding",
    }
    first_data_row = header_row + 1
    for row_offset, record in enumerate(records):
        row = first_data_row + row_offset
        values = {
            1: record.company_name,
            2: record.ticker,
            3: record.peer_category,
            4: record.cik,
            5: record.fiscal_year,
            6: record.filing_date.isoformat() if record.filing_date else None,
            7: record.status,
            8: record.revenue_mm,
            9: record.gross_profit_mm,
            11: record.operating_income_mm,
            12: record.depreciation_amortization_mm,
            15: record.cash_mm,
            16: record.debt_mm,
            18: record.shares_outstanding_mm,
            19: None,
        }
        for col_index, value in values.items():
            cell = ws.cell(row=row, column=col_index, value=value)
            cell.font = input_font if col_index >= 8 or col_index in {4, 5, 6, 7} else formula_font
            if col_index in metric_sources:
                source = _source_for_metric(record, metric_sources[col_index])
                cell.comment = Comment(
                    _comment_text(source, "Unavailable from SEC XBRL for this filing."),
                    "ADK comps agent",
                )

        formula_cells = {
            10: f'=IFERROR(I{row}/H{row},"")',
            13: f'=IF(OR(K{row}="",L{row}=""),"",K{row}+L{row})',
            14: f'=IFERROR(M{row}/H{row},"")',
            17: f'=IF(OR(P{row}="",O{row}=""),"",P{row}-O{row})',
            20: f'=IF(OR(R{row}="",S{row}=""),"",R{row}*S{row})',
            21: f'=IF(OR(T{row}="",Q{row}=""),"",T{row}+Q{row})',
            22: f'=IFERROR(U{row}/H{row},"")',
            23: f'=IFERROR(U{row}/M{row},"")',
        }
        for col_index, formula in formula_cells.items():
            cell = ws.cell(row=row, column=col_index, value=formula)
            cell.font = formula_font

    stats_start = first_data_row + len(records) + 1
    stat_labels = [
        ("Maximum", "MAX"),
        ("75th Percentile", "QUARTILE.INC", 3),
        ("Median", "MEDIAN"),
        ("25th Percentile", "QUARTILE.INC", 1),
        ("Minimum", "MIN"),
    ]
    stat_columns = [10, 14, 22, 23]
    data_last_row = first_data_row + len(records) - 1
    for offset, stat in enumerate(stat_labels):
        row = stats_start + offset
        ws.cell(row=row, column=1, value=stat[0]).font = header_font
        for col_index in range(1, len(headers) + 1):
            ws.cell(row=row, column=col_index).fill = stat_fill
        for col_index in stat_columns:
            col_letter = ws.cell(row=header_row, column=col_index).column_letter
            data_range = f"{col_letter}{first_data_row}:{col_letter}{data_last_row}"
            formula = (
                f'=IFERROR({stat[1]}({data_range},{stat[2]}),"")'
                if len(stat) == 3
                else f'=IFERROR({stat[1]}({data_range}),"")'
            )
            ws.cell(row=row, column=col_index, value=formula)

    for row in ws.iter_rows(min_row=first_data_row, max_row=stats_start + len(stat_labels), min_col=8, max_col=23):
        for cell in row:
            if cell.column in {10, 14}:
                cell.number_format = "0.0%"
            elif cell.column in {22, 23}:
                cell.number_format = "0.0x"
            elif cell.column == 19:
                cell.number_format = "$0.00"
            else:
                cell.number_format = '#,##0;(#,##0);"-"'
            cell.alignment = Alignment(horizontal="center", vertical="center")

    for col_index in range(1, len(headers) + 1):
        ws.column_dimensions[ws.cell(row=header_row, column=col_index).column_letter].width = 16
    ws.column_dimensions["A"].width = 30
    ws.freeze_panes = "A6"

    sources_headers = [
        "Company",
        "Ticker",
        "Metric",
        "Value",
        "Unit",
        "Source",
        "Confidence",
        "Form",
        "Fiscal Year",
        "Filing Date",
        "Accession Number",
        "XBRL Tag",
        "Notes",
    ]
    for col_index, header in enumerate(sources_headers, start=1):
        cell = sources_ws.cell(row=1, column=col_index, value=header)
        cell.fill = dark_fill
        cell.font = section_font
    source_row = 2
    for record in records:
        for source in record.sources:
            row_values = [
                record.company_name,
                record.ticker,
                source.get("metric"),
                source.get("value"),
                source.get("unit"),
                source.get("source"),
                source.get("confidence"),
                source.get("form"),
                source.get("fiscal_year"),
                source.get("filing_date"),
                source.get("accession_number"),
                source.get("xbrl_tag"),
                source.get("notes"),
            ]
            for col_index, value in enumerate(row_values, start=1):
                sources_ws.cell(row=source_row, column=col_index, value=value)
            source_row += 1
    for col_index in range(1, len(sources_headers) + 1):
        sources_ws.column_dimensions[sources_ws.cell(row=1, column=col_index).column_letter].width = 18
    sources_ws.column_dimensions["A"].width = 30
    sources_ws.column_dimensions["M"].width = 48

    review_ws["A1"] = "Peer Review"
    review_ws["A1"].font = Font(name="Times New Roman", size=12, bold=True)
    review_ws["A3"] = "Reviewer"
    review_ws["B3"] = review.decided_by_name or review.decided_by
    review_ws["A4"] = "Comment"
    review_ws["B4"] = review.comment
    review_ws["A6"] = "Selected Peer"
    review_ws["B6"] = "Ticker"
    review_ws["C6"] = "Category"
    review_ws["D6"] = "Rationale"
    for col_index in range(1, 5):
        review_ws.cell(row=6, column=col_index).fill = header_fill
        review_ws.cell(row=6, column=col_index).font = header_font
    for offset, peer in enumerate(review.selected_peers, start=7):
        review_ws.cell(row=offset, column=1, value=peer.company_name)
        review_ws.cell(row=offset, column=2, value=peer.ticker)
        review_ws.cell(row=offset, column=3, value=peer.peer_category)
        review_ws.cell(row=offset, column=4, value=peer.rationale)
    proposal_start = 9 + len(review.selected_peers)
    review_ws.cell(row=proposal_start, column=1, value="Original Suggested Peers")
    review_ws.cell(row=proposal_start, column=1).font = header_font
    for offset, peer in enumerate(proposal.suggested_peers, start=proposal_start + 1):
        review_ws.cell(row=offset, column=1, value=peer.company_name)
        review_ws.cell(row=offset, column=2, value=peer.ticker)
        review_ws.cell(row=offset, column=3, value=peer.peer_category)
        review_ws.cell(row=offset, column=4, value=peer.rationale)
    for col, width in {"A": 30, "B": 12, "C": 18, "D": 64}.items():
        review_ws.column_dimensions[col].width = width

    wb.save(path)
    return CompsExcelArtifact(
        path=str(path),
        filename=filename,
        mime_type=MIME_XLSX,
        generated_at=datetime.now(timezone.utc).isoformat(),
    )
